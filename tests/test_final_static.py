"""The Final Boss Start sheet: its own instructions, and the speedrun timer right under the hub (found in review)."""
import openpyxl

from excel_drills.content import FINAL, WORLDS
from excel_drills.course import Course
from excel_drills.final_builder import TIMER, build_final


def test_final_start_explains_rounds_and_keeps_the_timer_close(tmp_path):
    built = build_final(Course(WORLDS, FINAL), tmp_path)
    ws = openpyxl.load_workbook(built.path)["Start"]
    how_to = " ".join(str(ws[f"A{r}"].value or "") for r in range(5, 11))
    assert "BRIEFING" not in how_to and "WARM-UP" not in how_to
    assert "Round 1" in how_to and "GAUNTLET" in how_to and "START" in how_to
    replay = next(r for r in range(1, 60) if str(ws[f"A{r}"].value or "").startswith("Replay"))
    band = int(TIMER[0][1:]) - 1
    assert str(ws[f"A{band}"].value).startswith("SPEEDRUN") and band == replay + 2
