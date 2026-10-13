"""Static audits of built workbooks: accessibility, structure, metadata, allowed functions."""
import getpass
import re
import zipfile

import openpyxl
import pytest

from excel_drills.course import Course
from excel_drills.world_builder import build_world
from tests.demo_content import DEMO

BANNED = ("LAMBDA", "TEXTSPLIT", "TEXTBEFORE", "TEXTAFTER", "VSTACK", "HSTACK", "TAKE", "DROP", "CHOOSECOLS",
          "CHOOSEROWS", "TOCOL", "TOROW", "WRAPROWS", "WRAPCOLS", "EXPAND", "BYROW", "BYCOL", "MAP", "REDUCE",
          "SCAN", "MAKEARRAY", "GROUPBY", "PIVOTBY", "IMAGE", "RAND", "RANDBETWEEN", "RANDARRAY")


def _built(tmp_path_factory):
    from excel_drills.content import FINAL, WORLDS
    out = tmp_path_factory.mktemp("static")
    books = [build_world(Course((DEMO,)), DEMO, out)]
    real = Course(WORLDS, FINAL)
    books += [build_world(real, w, out) for w in WORLDS]
    return books


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    return _built(tmp_path_factory)


def test_visible_sheets_have_title_and_purpose_and_no_merges(built):
    for b in built:
        wb = openpyxl.load_workbook(b.path)
        for ws in wb.worksheets:
            assert not ws.merged_cells.ranges, f"{b.path.name}/{ws.title} has merged cells"
            if ws.sheet_state == "visible":
                assert ws["A1"].value and ws["A2"].value, f"{b.path.name}/{ws.title} lacks A1/A2"


def test_fonts_are_at_least_11pt(built):
    for b in built:
        wb = openpyxl.load_workbook(b.path)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        assert cell.font.sz is None or cell.font.sz >= 11, f"{ws.title}!{cell.coordinate}"


def test_protection_policy(built):
    for b in built:
        wb = openpyxl.load_workbook(b.path)
        for ws in wb.worksheets:
            if ws.title.endswith("Arena") or ws.title == "Rename Me":
                assert not ws.protection.sheet, ws.title
            elif ws.sheet_state == "visible":
                assert ws.protection.sheet and ws.protection.password is None, ws.title
        assert wb["Key"].sheet_state == "veryHidden"


def test_metadata_has_no_personal_information(built):
    user = getpass.getuser().lower()
    for b in built:
        with zipfile.ZipFile(b.path) as z:
            for name in z.namelist():
                text = z.read(name).decode("utf-8", "replace")
                low = text.lower()
                assert user not in low or user in ("excel", "drills"), f"{name} mentions the user name"
                assert "c:\\" not in low and "users\\" not in low and "abspath" not in low, name
            core = z.read("docProps/core.xml").decode()
            assert "<dc:creator>excel-drills</dc:creator>" in core
            assert "2026-10-12T00:00:00Z" in core


def test_only_excel_2021_functions(built):
    pattern = re.compile(r"\b(" + "|".join(BANNED) + r")\(")
    for b in built:
        for fc in b.formulas:
            outside = "".join(fc.formula.split('"')[0::2]).upper()
            assert not pattern.search(outside), f"{fc.sheet}!{fc.cell}: {fc.formula}"


def test_challenge_rows_have_explicit_heights(built):
    for b in built:
        wb = openpyxl.load_workbook(b.path)
        for placed in b.sheets:
            ws = wb[placed.sheet]
            for pc in placed.challenges:
                assert ws.row_dimensions[pc.row].height, f"{placed.sheet} row {pc.row}"
