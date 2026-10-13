import openpyxl

from excel_drills.arena_sheet import render_arena, render_extra
from excel_drills.checks import Filled, Pick, Recall, Spill, Value
from excel_drills.keysheet import KeySheet
from excel_drills.model import Arena, Formula
from excel_drills.refs import resolve
from excel_drills.writer import Book


def _text(cell):
    """openpyxl returns dynamic-array formulas as ArrayFormula objects; give their formula text."""
    return getattr(cell.value, "text", cell.value)


def test_key_sheet_layout(tmp_path):
    book = Book(tmp_path / "k.xlsx", title="t", subject="s", keywords="k")
    book.sheet("Start")                      # XlsxWriter never hides the first (active) sheet; Key is always last
    key = KeySheet(book, book.sheet("Key"))
    res = lambda t: resolve(t, {}, "Arena")
    n1 = key.add("A W1", "hint one", Recall(("CTRL+A", "=SUM")), res)
    n2 = key.add("A T1", "hint two", Pick(("A", "B", "C"), "B"), res)
    n3 = key.add("A C1", "hint three", Value("F6", 10, "=«@C6»*«@D6»"), res)
    n4 = key.add("A C2", "h", Filled(("C6", "C7"), (1, 2), ("=«@B6»", "=«@B7»")), res)
    n5 = key.add("A C3", "h", Spill("K6", ((1,), (2,)), "=SEQUENCE(2)"), res)
    book.close()
    assert n1 == {"hint": "Key!$B$2", "keyrow": "Key!$E$2:$F$2", "opts": "Key!$E$2:$F$2"}
    assert n2["key"] == "Key!$C$3" and n2["opts"] == "Key!$E$3:$G$3"
    assert n3["key"] == "Key!$C$4"
    assert n4["key1"] == "Key!$E$5" and n4["key2"] == "Key!$F$5"
    assert n5["spill"] == "Key!$E$2001"
    ws = openpyxl.load_workbook(tmp_path / "k.xlsx")["Key"]
    assert ws.sheet_state == "veryHidden"
    assert ws["F2"].value == "=SUM" and ws["F2"].data_type == "s"       # text, never a formula
    assert _text(ws["C4"]).startswith("=") and "Arena" in _text(ws["C4"])


def test_arena_render(tmp_path):
    book = Book(tmp_path / "a.xlsx", title="t", subject="s", keywords="k")
    stage = book.sheet("S1-1 Warp Zones")
    ws = book.sheet("S1-1 Arena")
    arena = Arena(cells={"B5": "HEAD", "B6": 1, "C6": Formula("=B6*2", 2), "D9": "=not a formula"},
                  pass_code="QUASAR", inputs=("F6:F7",), headers=("B5",), number_formats={"B6": "0.0"},
                  validations={"F7": ("Bronze", "Gold")}, names={"TREASURE": "B6"})
    render_arena(book, ws, arena, title="ARENA ▸ S1-1 WARP ZONES", back_sheet=stage.name)
    render_extra(book, book.sheet("Rename Me"), "Rename this sheet to Vault.")
    book.close()
    wb = openpyxl.load_workbook(tmp_path / "a.xlsx")
    a = wb["S1-1 Arena"]
    assert a["A1"].value == "ARENA ▸ S1-1 WARP ZONES" and a["B3"].value == "QUASAR"
    assert a["D9"].data_type == "s" and _text(a["C6"]).startswith("=")
    assert not a.protection.sheet
    assert a["F6"].fill.fgColor.rgb.endswith("FFF4CC")
    assert "TREASURE" in wb.defined_names
    assert wb["Rename Me"]["A1"].value == "RENAME ME"


from excel_drills import theme as T
from excel_drills.checks import Recall as _Recall
from excel_drills.model import Challenge, Fixture, Item, Stage, find, formula, self_check
from excel_drills.stage_sheet import ChallengeTable, render_stage


def test_stage_sheet(tmp_path):
    book = Book(tmp_path / "s.xlsx", title="t", subject="s", keywords="k")
    ws = book.sheet("S9-1 Demo")
    aws = book.sheet("S9-1 Arena")
    key = KeySheet(book, book.sheet("Key"))
    render_arena(book, aws, Arena(cells={"C6": 2, "D6": 5}, pass_code="P"), title="ARENA", back_sheet=ws.name)
    item = Item("S9.a", "shortcut", "Ctrl+A", "Select the data block.", "Press twice for more.",
                "Which keys select the data block?", ("Ctrl+A",), "Think A for All.")
    recall = Challenge("S9-1.R1", "RECALL", item.prompt, item.hint, _Recall(("CTRL+A",)),
                       (Fixture("type", "«ans»", "Ctrl+A"),), ((Fixture("type", "«ans»", "Ctrl+Q"),),), ("S9.a",))
    missions = (
        find("S9-1.M1", "Type the word.", "Look down.", "COMET", items=("S9.a",), wrong=("STAR",)),
        formula("S9-1.M2", "Multiply C6 by D6 in F6.", "Use *.", "F6", "=C6*D6", 10, items=("S9.a",), wrong=("=C6+D6",)),
        self_check("S9-1.M3", "Bold B6, then pick Done.", "Ctrl+B.", items=("S9.a",)),
    )
    stage = Stage("S9-1", "Demo", "test the table", (item,), missions, None)
    placed = render_stage(book, ws, key, stage, warmups=[], training=[], recalls=[recall], arena_sheet=aws.name,
                          next_label="S9-2 Next", extra_names={})
    book.close()
    assert placed.max_xp == 40 and len(placed.challenges) == 4
    s = openpyxl.load_workbook(tmp_path / "s.xlsx")["S9-1 Demo"]
    assert s.freeze_panes == "A5" and s.protection.sheet
    assert s["A1"].value == "STAGE S9-1 ▸ DEMO" and s["A2"].value == "Goal: test the table"
    rows = {pc.label: pc.row for pc in placed.challenges}
    fm = rows["C3"]
    assert s[f"C{fm}"].hyperlink.location == "'S9-1 Arena'!F6"
    assert _text(s[f"D{fm}"]).startswith("=") and s.row_dimensions[fm].height > 15
    assert s[f"C{rows['C1']}"].number_format == "@"
    sqrefs = " ".join(str(dv.sqref) for dv in s.data_validations.dataValidation)
    assert f"C{rows['C1']}" in sqrefs and f"E{rows['C1']}" in sqrefs
    assert len(s.conditional_formatting) >= 1
