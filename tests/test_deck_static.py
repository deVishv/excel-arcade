import zipfile

import openpyxl

from excel_drills.content import FINAL, WORLDS
from excel_drills.course import Course
from excel_drills.deck_builder import build_deck


def test_deck_structure(tmp_path):
    course = Course(WORLDS, FINAL)
    deck = build_deck(course, tmp_path)
    wb = openpyxl.load_workbook(deck.path)
    ws = wb["Deck"]
    last = 4 + len(course.items)
    assert len(deck.cards) == len(course.items) and deck.cards[0].row == 5
    assert ws.auto_filter.ref == f"A4:K{last}"
    assert ws.protection.sheet and ws.protection.autoFilter is False      # filtering allowed
    assert ws["E5"].protection.locked is False and ws["F5"].protection.locked is True
    assert wb["Key"].sheet_state == "veryHidden"
    with zipfile.ZipFile(deck.path) as z:                                  # chart alt text for screen readers
        drawing = z.read("xl/drawings/drawing1.xml").decode("utf-8")
    assert 'descr="Column chart: how many cards are New' in drawing
